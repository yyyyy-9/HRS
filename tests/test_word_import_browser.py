"""WORD-IMPORT-UI-01: native-selected DOCX through the real Bridge and page.

Only the system file/folder selection is substituted. The shipped component,
scanner, recognition, OMML adapter, serializers and project API all execute.
"""
from pathlib import Path
import hashlib,json,os,unittest
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
from browser_support import ManualService,source,value
import studio


def native_link(p,label,target=None,anchor=None,bold=False):
    node=OxmlElement('w:hyperlink')
    if target is not None:node.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    if anchor is not None:node.set(qn('w:anchor'),anchor)
    for part in label if isinstance(label,list) else [label]:
        run=OxmlElement('w:r')
        if bold:
            properties=OxmlElement('w:rPr');properties.append(OxmlElement('w:b'));run.append(properties)
        text=OxmlElement('w:t');text.text=part;run.append(text);node.append(run)
    p._p.append(node)
    return node


def native_bookmark(p,name,ident):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(ident));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(ident));p._p.insert(0,start);p._p.append(end)


class WordImport(ManualService):
    def setUp(self):
        super().setUp()
        self.input=self.folder/(self._testMethodName+'.docx')
        self.png=self.folder/(self._testMethodName+'.png')
        Image.new('RGB',(48,24),'navy').save(self.png)
        doc=Document();doc.add_paragraph('导入报告',style='Title');doc.add_heading('目的',level=1)
        doc.add_paragraph('保留量值 12.34 nm，普通 @fig:literal 不激活。')
        doc.add_paragraph('表1 参数表')
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='参数';table.cell(0,1).text='数值'
        table.cell(1,0).text='行内公式 '
        math=OxmlElement('m:oMath');run=OxmlElement('m:r');text=OxmlElement('m:t');text.text='x'
        run.append(text);math.append(run);table.cell(1,0).paragraphs[0]._p.append(math)
        table.cell(1,1).text='12.34'
        p=doc.add_paragraph();p.add_run().add_picture(str(self.png));p.add_run('   ');p.add_run().add_picture(str(self.png))
        doc.add_paragraph('图1 两幅原图')
        doc.save(self.input);self.original=self.input.read_bytes()
        self.picker=patch.object(studio,'pick_import_file',return_value=self.input)
        self.picker.start();self.addCleanup(self.picker.stop)

    def read_word(self):
        self.page.locator('#menu-open>summary').click();self.page.locator('#btn-import-word').click()
        self.page.wait_for_function("document.querySelector('#modal-title').firstChild.textContent==='导入 Word'")
        self.assertEqual(self.page.locator('#modal-title .beta-badge').inner_text(),'beta')
        self.page.locator('#modal-ok').click()
        self.page.wait_for_function("document.querySelector('#modal-title').firstChild.textContent==='核对 Word 导入结果'||document.querySelector('#toast').textContent.includes('导入失败')",timeout=90000)
        self.assertEqual(self.page.locator('#modal-title').evaluate('n=>n.firstChild.textContent'),'核对 Word 导入结果',self.page.locator('#toast').inner_text())
        self.assertEqual(self.page.locator('#modal-title .beta-badge').inner_text(),'beta')

    def save_import(self,name):
        self.choose_destination(name)
        self.page.wait_for_function("document.querySelector('#toast').textContent.includes('Word 已导入')||document.querySelector('#toast').textContent.includes('导入失败')",timeout=30000)

    def choose_destination(self,name):
        self.page.locator('#modal-ok').click()
        self.page.wait_for_selector('#modal-body input[aria-label="项目文件夹名称"]')
        self.page.locator('#modal-body input').fill(name);self.page.locator('#modal-ok').click()
        self.page.locator('.confirm-window button').filter(has_text='确认位置').click()

    def test_real_docx_table_math_group_save_reopen_and_archive_without_node(self):
        # The page's browser is already running; GUI import must never launch
        # another one or depend on a system Node executable.
        with patch.dict(os.environ,{'HRS_BRIDGE_NODE':'Z:/absent-node.exe','CHROMIUM_PATH':'Z:/absent-browser.exe'}):
            self.read_word()
            self.assertIn('结构转换完成',self.page.locator('#word-import-summary').inner_text())
            self.save_import('Word导入专项')
        target=self.destination/'Word导入专项';self.assertTrue(target.is_dir(),self.page.locator('#toast').inner_text())
        md=value(self.page);self.assertIn('12.34 nm',md);self.assertIn('\\@fig:literal',md)
        self.assertIn('$x$',md);self.assertIn('report-table-v1',md)
        model=studio.analyze_report_contract(md,assets=target/'assets')
        self.assertEqual(len([b for b in model['blocks'] if b['type']=='image-group']),1)
        self.page.wait_for_function("()=>{const a=[...document.querySelectorAll('#preview-content img')];a.forEach(i=>i.loading='eager');return a.length===2&&a.every(i=>i.naturalWidth>0)}")
        self.assertEqual((target/'report.md').read_bytes(),md.encode('utf8'))
        assets=list((target/'assets').glob('*.png'));self.assertEqual(len(assets),1)
        self.assertEqual(assets[0].read_bytes(),self.png.read_bytes())
        receipt=json.loads((target/'word-import.json').read_text('utf8'))
        self.assertEqual(receipt['source_sha256'],hashlib.sha256(self.original).hexdigest())
        self.assertEqual(receipt['status'],'COMPLETE');self.assertEqual(receipt['recognition']['model_calls'],0)
        self.assertEqual(self.input.read_bytes(),self.original)
        html=self.page.evaluate("()=>ReportStudioProgrammatic.exportHTML({sourceMode:'recompilable'})")
        studio.audit_export(html);self.assertIn('report-source',html)
        self.page.reload();self.page.evaluate('()=>ReportStudioProgrammatic.ready()')
        self.page.locator('#recent-list button').filter(has_text='导入报告').first.click()
        self.page.wait_for_function("document.querySelector('#markdown').sourceEditor.value.includes('12.34 nm')")
        self.assertEqual(value(self.page),md);self.assertFalse(self.server.word_import.jobs)
        if self.evidence:self.page.screenshot(path=str(self.evidence/'imported.png'))

    def test_cancel_review_keeps_current_source_and_draft(self):
        source(self.page,'# 当前未保存稿\n\n保留😀\n');before=value(self.page)
        self.read_word();self.page.locator('#modal-cancel').click()
        self.page.wait_for_function("!document.querySelector('#app').inert")
        self.assertEqual(value(self.page),before);self.assertEqual(self.input.read_bytes(),self.original)
        self.assertFalse(self.server.word_import.jobs)

    def test_changed_original_is_rejected_before_project_commit(self):
        before=value(self.page);self.read_word()
        self.input.write_bytes(self.original+b'changed')
        self.save_import('来源已变不可导入')
        self.assertIn('已变化',self.page.locator('#toast').inner_text())
        self.assertEqual(value(self.page),before);self.assertFalse((self.destination/'来源已变不可导入').exists())
        self.assertFalse(self.server.word_import.jobs)

    def test_unsupported_nested_table_keeps_current_project(self):
        doc=Document();cell=doc.add_table(rows=1,cols=1).cell(0,0);cell.add_table(rows=1,cols=1);doc.save(self.input)
        before=value(self.page)
        self.page.locator('#menu-open>summary').click();self.page.locator('#btn-import-word').click();self.page.locator('#modal-ok').click()
        self.page.wait_for_function("document.querySelector('#toast').textContent.includes('导入失败')",timeout=90000)
        self.assertEqual(value(self.page),before);self.assertFalse(self.server.word_import.jobs)

    def test_needs_review_keeps_unconfirmed_prose_and_lists_it(self):
        text='这是被误用标题样式的一段完整正文，需要保留全部文字而不能强制认作章节。'*6
        doc=Document();doc.add_paragraph('待确认样本',style='Title');doc.add_heading(text,level=1);doc.save(self.input)
        self.original=self.input.read_bytes();self.read_word()
        self.assertIn('待确认',self.page.locator('#word-import-summary').inner_text())
        self.assertIn(text,self.page.locator('#word-import-diagnostics').inner_text())
        self.save_import('Word待确认项')
        md=value(self.page);self.assertIn(text,md);self.assertNotIn('## '+text,md)
        receipt=json.loads((self.destination/'Word待确认项/word-import.json').read_text('utf8'))
        self.assertEqual(receipt['status'],'NEEDS_REVIEW');self.assertEqual(receipt['recognition']['unresolved'],1)
        self.assertIn('待确认',self.page.locator('#toast').inner_text())

    def test_missing_formula_tool_and_runtime_leave_original_unchanged(self):
        before=value(self.page)
        with patch.object(self.server.word_export,'pandoc',return_value=None):
            self.page.locator('#menu-open>summary').click();self.page.locator('#btn-import-word').click();self.page.locator('#modal-ok').click()
            self.page.wait_for_function("document.querySelector('#toast').textContent.includes('导入失败')",timeout=30000)
        self.assertIn('Pandoc',self.page.locator('#toast').inner_text());self.assertEqual(value(self.page),before)
        with patch.object(self.server.word_export,'runtime',return_value=None):
            self.page.locator('#menu-open>summary').click();self.page.locator('#btn-import-word').click();self.page.locator('#modal-ok').click()
            self.assertTrue(self.page.locator('#modal-layer').is_visible())
            self.assertIn('请先选择',self.page.locator('#word-import-runtime-state').inner_text())
            self.page.locator('#modal-cancel').click()
        self.assertFalse(self.server.word_import.jobs);self.assertEqual(self.input.read_bytes(),self.original)

    def test_cancel_waiting_serializer_stops_owned_process_and_releases_component(self):
        before=value(self.page)
        result=self.page.evaluate("""async()=>{const r=await fetch('/api/import/word/start',{method:'POST',headers:{'Content-Type':'application/json','X-HRS-Project-Action':'user-action'},body:'{}'});return r.json()}""")
        self.assertTrue(result['ok'],result);token=result['job_token']
        pending=self.page.evaluate("""async token=>{for(let i=0;i<100;i++){const r=await fetch('/api/import/word/step',{method:'POST',headers:{'Content-Type':'application/json','X-HRS-Project-Action':'user-action'},body:JSON.stringify({job_token:token})});const v=await r.json();if(v.state==='page'||v.state==='error')return v;await new Promise(r=>setTimeout(r,50));}throw Error('No serializer request')}""",token)
        self.assertEqual(pending['state'],'page',pending)
        job=self.server.word_import.jobs[token];process=job['process'];folder=Path(job['temporary'].name)
        stale=self.page.evaluate("""async token=>{const r=await fetch('/api/import/word/step',{method:'POST',headers:{'Content-Type':'application/json','X-HRS-Project-Action':'user-action'},body:JSON.stringify({job_token:token,seq:99999,result:[]})});return {status:r.status,...await r.json()}}""",token)
        self.assertEqual(stale['status'],400);self.assertIsNone(process.poll())
        cancelled=self.page.evaluate("""async token=>{const r=await fetch('/api/import/word/cancel',{method:'POST',headers:{'Content-Type':'application/json','X-HRS-Project-Action':'user-action'},body:JSON.stringify({job_token:token})});return r.json()}""",token)
        self.assertTrue(cancelled['ok'],cancelled);self.assertIsNotNone(process.poll())
        self.assertFalse(folder.exists());self.assertFalse(self.server.word_export.running.locked())
        self.assertEqual(value(self.page),before);self.assertFalse(self.server.word_import.jobs)

    def test_registry_write_failure_keeps_verified_new_project_and_current_source(self):
        before=value(self.page);self.read_word()
        with patch.object(self.server.registry,'register',side_effect=OSError('test index lock')):
            self.save_import('登记失败保留工程')
        target=self.destination/'登记失败保留工程'
        self.assertIn('工程已完整保存',self.page.locator('#toast').inner_text())
        self.assertEqual(value(self.page),before);self.assertTrue((target/'assets').is_dir())
        self.assertIn('12.34', (target/'report.md').read_text('utf8'))
        self.assertFalse(self.server.word_import.jobs)
        with patch.object(studio,'pick_project_parent',return_value=target):
            self.page.locator('#menu-open>summary').click();self.page.locator('#btn-project-open').click()
            self.page.wait_for_function("document.querySelector('#markdown').sourceEditor.value.includes('12.34')")
        self.assertEqual(self.input.read_bytes(),self.original)

    def test_cancel_switch_after_success_preserves_dirty_current_draft(self):
        text='# 当前未保存稿\n\n请保留源稿和草稿😀。\n';source(self.page,text)
        self.read_word();self.choose_destination('导入成功但暂不切换')
        self.page.wait_for_function("document.querySelector('#modal-title').textContent==='有未保存的修改'")
        self.page.locator('#modal-cancel').click()
        self.page.wait_for_function("document.querySelector('#toast').textContent.includes('当前文稿保持')")
        self.assertEqual(value(self.page),text);self.assertEqual((self.project/'report.md').read_bytes(),b'# Initial\n')
        self.assertTrue((self.destination/'导入成功但暂不切换/report.md').is_file())
        self.assertFalse(self.server.word_import.jobs)

    def test_own_export_uses_companion_mapping_for_exact_unchanged_source(self):
        from test_word_export_browser import WordUI
        self.docid='f'*32;self.server.registry.register(self.docid,self.project)
        self.page.evaluate('id=>ReportStudioProgrammatic.open(id)',self.docid)
        source(self.page,'# 自身导出回编\n\n## 目的\n\n**强调** 与 $x^2$，中文😀。\n\n| 参数 | 数值 |\n| --- | --- |\n| 长度 | 12.34 |\n')
        self.page.locator('#btn-save').click();self.page.wait_for_function("document.querySelector('#save-indicator').textContent.trim()==='已保存'")
        expected=value(self.page);WordUI.open_word(self);output=WordUI.compile(self)
        raw=output.read_bytes();companion=Path(str(output)+'.hrs-bridge.json');mapping=companion.read_bytes()
        with patch.object(studio,'pick_import_file',return_value=output):
            self.read_word();self.save_import('自身Word恢复')
        restored=self.destination/'自身Word恢复';self.assertEqual((restored/'report.md').read_bytes(),expected.encode('utf8'))
        receipt=json.loads((restored/'word-import.json').read_text('utf8'))
        self.assertEqual(receipt['mode'],'provenance-aware');self.assertEqual(output.read_bytes(),raw)
        self.assertEqual(companion.read_bytes(),mapping);self.assertFalse(self.server.word_import.jobs)

    def test_hyperlinks_beta_preview_html_and_native_word_keep_real_destinations(self):
        """WORD-HYPERLINK-01: labels/OMML/runs and forward targets survive the actual UI."""
        doc=Document();doc.add_paragraph('链接完整链',style='Title')
        p=doc.add_paragraph('参考');native_link(p,['外部','资料'],'https://example.org/data?x=1&y=2',bold=True);p.add_run('，请');native_link(p,'跳至目的',anchor='_Purpose');p.add_run('或');native_link(p,'跳至结果',anchor='_Results');p.add_run('。')
        p=doc.add_heading('目的',level=1);native_bookmark(p,'_Purpose',50)
        p=doc.add_paragraph();native_link(p,'表1 参数表','https://example.org/table')
        table=doc.add_table(rows=1,cols=1);node=native_link(table.cell(0,0).paragraphs[0],'公式 ','https://example.org/math')
        math=OxmlElement('m:oMath');run=OxmlElement('m:r');text=OxmlElement('m:t');text.text='x';run.append(text);math.append(run);node.append(math)
        p=doc.add_heading('结果',level=1);native_bookmark(p,'_Results',51);doc.add_paragraph('结果正文保持。')
        doc.save(self.input);self.original=self.input.read_bytes()
        for theme in ('blue','warm','navy'):
            self.page.locator('#color-theme').select_option(theme)
            self.page.wait_for_function('t=>document.documentElement.dataset.colorTheme===t',arg=theme)
            self.page.locator('#menu-open>summary').click()
            badge=self.page.locator('#btn-import-word .beta-badge');self.assertEqual(badge.inner_text(),'beta')
            self.assertEqual(badge.evaluate('e=>getComputedStyle(e).color'),self.page.locator('#btn-export-word .beta-badge').evaluate('e=>getComputedStyle(e).color'))
            metric=badge.evaluate('e=>{const r=e.getBoundingClientRect(),b=e.parentElement.getBoundingClientRect();return {right:r.right,top:r.top,parentTop:b.top,width:innerWidth}}')
            self.assertLess(metric['top'],metric['parentTop']+8);self.assertLessEqual(metric['right'],metric['width'])
            self.page.locator('#menu-open>summary').click()
        self.read_word();self.save_import('链接完整链')
        project=self.destination/'链接完整链';md=value(self.page)
        model=studio.analyze_report_contract(md,assets=project/'assets')
        ids={b['text']:b['id'] for b in model['blocks'] if b['type']=='heading' and b.get('id')}
        href='#object-'+ids['目的'];target=self.page.locator('#object-'+ids['目的'])
        self.assertEqual(target.count(),1);self.page.locator('#preview-content a').filter(has_text='跳至目的').click()
        self.page.wait_for_function('href=>location.hash===href',arg=href)
        self.assertEqual(self.page.locator('#preview-content a').filter(has_text='外部资料').get_attribute('href'),'https://example.org/data?x=1&y=2')
        self.assertEqual(self.page.locator('#preview-content a').filter(has_text='外部资料').locator('strong').inner_text(),'外部资料')
        self.assertIn('[公式 $x$](<https://example.org/math>)',md)
        html=self.page.evaluate("()=>ReportStudioProgrammatic.exportHTML({sourceMode:'recompilable'})");studio.audit_export(html)
        offline=self.folder/'links.html';offline.write_text(html,encoding='utf8')
        reader=self.page.context.new_page()
        try:
            reader.goto(offline.as_uri());reader.locator('a').filter(has_text='跳至结果').first.click()
            reader.wait_for_function('id=>location.hash==="#object-"+id',arg=ids['结果'])
        finally:reader.close()
        from test_word_export_browser import WordUI
        old_project=self.project;self.project=project
        try:WordUI.open_word(self);output=WordUI.compile(self)
        finally:self.project=old_project
        word=Document(output);links=word._element.xpath('//w:hyperlink');names=set(word._element.xpath('//w:bookmarkStart/@w:name'))
        internal={''.join(n.xpath('.//w:t/text()')):n.get(qn('w:anchor')) for n in links if n.get(qn('w:anchor'))}
        self.assertIn(internal['跳至目的'],names);self.assertIn(internal['跳至结果'],names)
        destinations={word.part.rels[n.get(qn('r:id'))].target_ref for n in links if n.get(qn('r:id'))}
        self.assertIn('https://example.org/data?x=1&y=2',destinations)
        self.assertEqual(len(word._element.xpath('//w:hyperlink//m:oMath')),1)
        # Changing only the target must be read from Word, never restored from
        # the old Markdown snapshot just because the visible text is unchanged.
        next(n for n in links if ''.join(n.xpath('.//w:t/text()'))=='跳至目的').set(qn('w:anchor'),internal['跳至结果'])
        word.save(output)
        with patch.object(studio,'pick_import_file',return_value=output):
            self.read_word();self.save_import('修改链接目标后回编')
        self.assertIn('[跳至目的](<#object-'+ids['结果']+'>)',value(self.page))
        self.assertEqual(self.input.read_bytes(),self.original)
        if self.evidence:self.page.screenshot(path=str(self.evidence/'links.png'))

    def test_unknown_and_unsafe_link_targets_are_shown_and_saved_for_review(self):
        doc=Document();doc.add_paragraph('待确认链接',style='Title');p=doc.add_paragraph('参见')
        native_link(p,'未知内部目标',anchor='_Unknown');p.add_run('、');native_link(p,'危险目标','javascript:alert(1)');p.add_run('、');native_link(p,'相对原目标','../docs/paper.pdf');p.add_run('。')
        doc.save(self.input);self.original=self.input.read_bytes();self.read_word()
        details=self.page.locator('#word-import-diagnostics').inner_text()
        for text in ('#_Unknown','javascript:alert(1)','../docs/paper.pdf'):self.assertIn(text,details)
        self.save_import('待确认链接');md=value(self.page)
        self.assertIn('未知内部目标、危险目标、相对原目标。',md);self.assertNotIn('](',md)
        receipt=json.loads((self.destination/'待确认链接/word-import.json').read_text('utf8'))
        self.assertEqual(receipt['status'],'NEEDS_REVIEW');self.assertEqual(receipt['recognition']['unresolved'],3)
        self.assertEqual(self.input.read_bytes(),self.original)


if __name__=='__main__':unittest.main()

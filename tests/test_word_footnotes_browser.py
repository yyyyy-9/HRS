"""FN-01/TABLE-02: table notes, true Word notes and one-action alignment."""
from pathlib import Path
import base64, json, os, unittest, zipfile
from lxml import etree as E
from docx import Document
from browser_support import ManualService
from browser_support import source, value
import test_word_export_browser as previous
import studio

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
TABLE={'schema':'report-table-v1','id':'tbl-notes','caption':'增强参数表[^caption]',
       'rows':[[{'text':'参数','rowspan':1,'colspan':1},{'text':'说明','rowspan':1,'colspan':1}],
               [{'text':'增强表[^shared]','rowspan':1,'colspan':1},{'text':'**加粗**与==重点==','rowspan':1,'colspan':1}],
               [{'text':'','rowspan':1,'colspan':2},None]]}
TEXT=('# 脚注与表格\n\n## 说明\n\n正文首次引用[^body]，再次引用[^shared]。\n\n'
      '| 参数 | 说明 |\n| --- | --- |\n| 普通表[^shared] | **加粗** |\n: 普通参数表[^caption]\n\n'
      +'```report-table\n'+json.dumps(TABLE,ensure_ascii=False)+'\n```\n\n'
      +'[^body]: 正文脚注 **加粗说明**。\n\n[^shared]: 表格共用脚注与[链接](https://example.com/)。\n'
      +'    第二段仍属于脚注。\n\n[^caption]: 表题脚注说明。\n\n`[^literal]`\n\n```md\n[^code]: 代码示例\n```\n')


class FootnotesTables(ManualService):
    open_word=previous.WordUI.open_word
    compile=previous.WordUI.compile

    def setUp(self):
        super().setUp();self.docid='f'*32;self.server.registry.register(self.docid,self.project)
        self.page.evaluate('id=>ReportStudioProgrammatic.open(id)',self.docid)
        self.page.set_default_timeout(20000)

    def test_01_table_refs_share_document_numbering_and_offline_navigation(self):
        """Both table kinds resolve against document definitions; source stays exact."""
        source(self.page,TEXT)
        self.page.wait_for_function("document.querySelectorAll('#preview-content table').length===2")
        refs=self.page.locator('#preview-content a[data-hrs-footnote-ref]')
        self.assertEqual(refs.all_text_contents(),['1','2','3','2','3','2'])
        self.assertEqual(self.page.locator('#preview-content td a.hrs-note-link').count(),2)
        self.assertEqual(self.page.locator('#preview-content .hrs-footnotes li').count(),3)
        self.assertEqual(self.page.locator('#preview-content .hrs-table-kind a.hrs-note-link').count(),2)
        self.page.locator('#preview-content .hrs-table-kind a.hrs-note-link').first.click()
        self.assertTrue(self.page.evaluate("document.activeElement.id==='hrs-note-3'"))
        self.page.locator('#preview-content td a.hrs-note-link').first.click()
        self.assertTrue(self.page.evaluate("document.activeElement.id==='hrs-note-2'"))
        self.page.locator('#hrs-note-2 .hrs-note-back').last.click()
        self.assertEqual(value(self.page),TEXT)
        self.page.locator('#render-profile').select_option('engineering')
        self.page.wait_for_function("document.querySelectorAll('#preview-content .hrs-table-heading-caption a.hrs-note-link').length===2")
        self.page.locator('#preview-content .hrs-table-heading-caption a.hrs-note-link').last.click()
        self.assertTrue(self.page.evaluate("document.activeElement.id==='hrs-note-3'"))
        self.page.locator('#render-profile').select_option('free')
        self.page.wait_for_function("document.querySelector('#preview-content .hrs-table-mini')!==null")
        html=self.page.evaluate("()=>ReportStudioProgrammatic.exportHTML({sourceMode:'recompilable'})")
        studio.audit_export(html)
        self.assertIn('1275a6830b70fb59cf07ac16be15e2dc5ce9c87b2af2345db2c3d2bf27dbeee8',studio.LEGACY_READER_SHA256)
        path=self.folder/'notes.html';path.write_bytes(html.encode('utf8'))
        reader=self.browser.new_page();reader.goto(path.as_uri())
        self.assertEqual(reader.locator('main td a.hrs-note-link').count(),2)
        reader.locator('main td a.hrs-note-link').first.click()
        self.assertEqual(reader.locator('main .hrs-footnotes li').count(),3)
        self.assertEqual(reader.locator('main .hrs-table-kind a.hrs-note-link').count(),2)
        reader.close()
        if self.evidence:self.page.screenshot(path=str(self.evidence/'tables-notes.png'))

    def test_02_all_columns_center_is_one_undoable_edit_for_both_table_kinds(self):
        """A single explicit action aligns every column, including empty/merged cells."""
        source(self.page,TEXT)
        self.page.locator('#preview-content .report-table-wrap').first.locator('td').first.dblclick()
        self.page.get_by_role('button',name='所有列水平居中',exact=True).click()
        self.page.locator('#modal-cancel').click();self.assertEqual(value(self.page),TEXT)
        for index in (0,1):
            table=self.page.locator('#preview-content .report-table-wrap').nth(index)
            table.locator('td').first.dblclick()
            self.page.get_by_role('button',name='所有列水平居中',exact=True).click()
            self.page.get_by_role('button',name='撤销表格',exact=True).click()
            self.page.get_by_role('button',name='重做表格',exact=True).click()
            self.page.locator('#modal-ok').click()
            self.page.wait_for_function("i=>[...document.querySelectorAll('#preview-content .report-table-wrap')][i].querySelectorAll('td,th').length>0 && [...[...document.querySelectorAll('#preview-content .report-table-wrap')][i].querySelectorAll('td,th')].every(e=>getComputedStyle(e).textAlign==='center')",arg=index)
        final=value(self.page);self.assertIn('| :---: | :---: |',final)
        self.assertIn('"center"',final)
        self.assertIn('[^shared]',final)
        self.open_word();output=self.compile()
        for table in Document(output).tables:
            self.assertTrue(all(p.alignment==1 for row in table.rows for cell in row.cells for p in cell.paragraphs))

    def test_03_word_has_native_notes_from_body_and_tables_and_exact_source_return(self):
        """Every marker reaches footnotes.xml, with preserved code and rich note text."""
        source(self.page,TEXT);self.open_word()
        # R25 defaults to stripping extra emphasis. This case explicitly tests
        # the supported rich-note policy; keep its original native assertions.
        self.page.locator('#word-keep-emphasis').check();output=self.compile()
        with zipfile.ZipFile(output) as z:
            body=E.fromstring(z.read('word/document.xml'))
            notes=E.fromstring(z.read('word/footnotes.xml'))
            refs=body.xpath('//w:footnoteReference/@w:id',namespaces=NS)
            self.assertEqual(len(refs),3)
            definitions=notes.xpath('w:footnote[not(@w:type)]',namespaces=NS)
            self.assertEqual(len(definitions),3)
            self.assertEqual(set(refs),{n.get('{'+NS['w']+'}id') for n in definitions})
            self.assertEqual(len(body.xpath('//w:fldSimple[contains(@w:instr,"NOTEREF")]',namespaces=NS)),3)
            self.assertEqual(len(body.xpath('//w:tbl//w:fldSimple[contains(@w:instr,"NOTEREF")]',namespaces=NS)),2)
            self.assertEqual(len(body.xpath('//w:fldSimple[contains(@w:instr,"NOTEREF")]//w:vertAlign[@w:val="superscript"]',namespaces=NS)),3)
            self.assertEqual(len(notes.xpath('//w:footnote[@w:type="separator"]',namespaces=NS)),1)
            self.assertTrue(notes.xpath('//w:b',namespaces=NS))
            self.assertTrue(notes.xpath('//w:hyperlink[@r:id]',namespaces={**NS,'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}))
            self.assertTrue(z.read('word/_rels/footnotes.xml.rels'))
        self.assertIn('[^literal]','\n'.join(p.text for p in Document(output).paragraphs))
        receipt=json.loads(Path(str(output)+'.hrs-bridge.json').read_text('utf8'))
        self.assertEqual(base64.b64decode(receipt['source_b64']).decode('utf8'),TEXT)
        self.assertEqual((self.project/'report.md').read_bytes(),b'# Initial\n')
        self.assertEqual(value(self.page),TEXT)


if __name__=='__main__':unittest.main()

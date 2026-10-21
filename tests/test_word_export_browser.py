"""Real HRS UI -> local snapshot -> offline Bridge/WPS -> DOCX.

Only native chooser return values use test-owned paths. Compilation, template
handling, draft upload, source checks and HTTP routes run without mocks.
"""
from pathlib import Path
import unittest
import hashlib,json,os,shutil,zipfile
from unittest.mock import patch
from docx import Document
from lxml import etree as E
from PIL import Image
from browser_support import ManualService
from browser_support import source,value
import studio

TEXT='# Word 界面专项\n\n## 试验目的\n\n当前工作快照：中文 😀 与 < > &。\n\n1. 第一项\n2. 第二项\n\n### 细节\n\n```python\nprint("keep")\n```\n\n| 参数 | 数值 |\n| --- | --- |\n| 功率 | 12 |\n'


class WordUI(ManualService):
    def setUp(self):
        super().setUp()
        self.docid=next((key for key,path in self.server.registry.mapping.items() if Path(path).resolve()==self.project.resolve()),'f'*32)
        self.server.registry.register(self.docid,self.project)
        self.page.evaluate('id=>ReportStudioProgrammatic.open(id)',self.docid)
        self.page.set_default_timeout(20000)

    def open_word(self):
        # This suite protects the engineering Word dialog/assembler. Free-note
        # preparation/resume is exercised once in NativeReportWord, not repeated
        # in every low-level template/notes case.
        if self.page.locator('#render-profile').input_value()!='engineering':
            self.page.locator('#render-profile').select_option('engineering')
            self.page.wait_for_function("document.body.dataset.renderProfile==='engineering'&&!document.querySelector('#render-profile').disabled")
        self.page.locator('#menu-export>summary').click();self.page.locator('#btn-export-word').click()
        working=self.page.locator('#modal-extra')
        if working.is_visible():self.page.locator('#modal-ok').click()
        self.page.wait_for_selector('#word-template')

    def compile(self):
        self.page.locator('#modal-ok').click()
        self.page.wait_for_function("document.querySelector('#modal-title').textContent.includes('Word 已生成') || document.querySelector('#word-export-status')?.textContent.startsWith('导出未完成')",timeout=90000)
        self.assertIn('Word 已生成',self.page.locator('#modal-title').inner_text(),self.page.locator('#word-export-status').all_inner_texts())
        output=max((self.project/'output').glob('*.docx'),key=lambda p:p.stat().st_mtime_ns)
        self.assertTrue(Path(str(output)+'.hrs-bridge.json').is_file())
        if self.evidence:
            shutil.copyfile(output,self.evidence/output.name)
            shutil.copyfile(Path(str(output)+'.hrs-bridge.json'),self.evidence/(output.name+'.hrs-bridge.json'))
        self.page.locator('#modal-ok').click()
        return output

    def test_01_themed_beta_menu_and_dialog(self):
        colors=[]
        for theme in ['blue','warm','navy']:
            self.page.locator('#color-theme').select_option(theme)
            self.page.wait_for_function('t=>document.documentElement.dataset.colorTheme===t',arg=theme)
            self.page.locator('#menu-export>summary').click()
            badge=self.page.locator('#btn-export-word .beta-badge')
            self.assertEqual(badge.inner_text(),'beta')
            metric=badge.evaluate('e=>{const r=e.getBoundingClientRect(),p=e.parentElement.getBoundingClientRect();return {color:getComputedStyle(e).color,border:getComputedStyle(e).borderColor,x:r.x,y:r.y,right:r.right,pY:p.y,pX:p.x,width:innerWidth}}')
            self.assertGreater(metric['x'],metric['pX']);self.assertLess(metric['y'],metric['pY']+8)
            self.assertLessEqual(metric['right'],metric['width']);self.assertEqual(metric['color'],metric['border']);colors.append(metric['color'])
            if self.evidence:
                self.page.screenshot(path=str(self.evidence/(theme+'-menu.png')))
                self.page.locator('#menu-export .top-menu-list').screenshot(path=str(self.evidence/(theme+'-menu-detail.png')))
            self.page.locator('#menu-export>summary').click()
        self.assertEqual(len(set(colors)),3)
        self.page.set_viewport_size({'width':980,'height':780});self.open_word()
        self.assertEqual(self.page.locator('#modal-title .beta-badge').inner_text(),'beta')
        self.assertEqual(self.page.locator('#word-runtime-state').get_attribute('data-ready'),'true')
        if self.evidence:self.page.screenshot(path=str(self.evidence/'word-dialog-narrow.png'))
        self.page.keyboard.press('Escape');self.assertFalse(self.page.locator('#modal-layer').is_visible())
        with patch.object(self.server.word_export,'runtime',return_value=None):
            self.open_word();self.assertEqual(self.page.locator('#word-runtime-state').get_attribute('data-ready'),'false')
            self.page.locator('#modal-ok').click();self.assertIn('Word',self.page.locator('#word-export-status').inner_text())
            self.page.keyboard.press('Escape');self.assertFalse(self.page.locator('#modal-layer').is_visible())

    def test_02_default_template_working_source_and_staged_image(self):
        original=(self.project/'report.md').read_bytes();source(self.page,TEXT)
        image=self.project.parent/'draft-image.png';Image.new('RGB',(48,24),(25,110,180)).save(image)
        self.page.locator('#file-image').set_input_files(str(image))
        self.page.wait_for_function("document.querySelector('#preview-content img')?.naturalWidth>0")
        md=value(self.page);self.assertIn('report-figure',md)
        self.open_word();self.page.locator('#word-default-template').click()
        if self.evidence:self.page.screenshot(path=str(self.evidence/'default-dialog.png'))
        output=self.compile();doc=Document(output)
        with patch('os.startfile',create=True) as opened:
            with self.page.expect_response(lambda r:'/api/word/open' in r.url) as response:
                self.page.locator('#toast').get_by_role('button',name='打开所在文件夹',exact=True).click()
            self.assertTrue(response.value.ok,response.value.text());opened.assert_called_once_with(str(output.parent))
        self.assertIn('Word 界面专项',[p.text for p in doc.paragraphs]);self.assertEqual(len(doc.inline_shapes),1)
        self.assertTrue(any('第一项' in p.text for p in doc.paragraphs));self.assertEqual(len(doc.tables),1)
        self.assertEqual((self.project/'report.md').read_bytes(),original)
        self.assertEqual(value(self.page),md)
        meta=json.loads(Path(str(output)+'.hrs-bridge.json').read_text(encoding='utf8'))
        self.assertEqual(meta['source_sha256'],hashlib.sha256(md.encode('utf8')).hexdigest())

    def test_03_cover_template_remembering_and_failure_protection(self):
        source(self.page,TEXT)
        runtime=self.server.word_export.runtime();self.assertIsNotNone(runtime)
        template=runtime['wps']/'templates/cover_template_v0311.docx'
        before=hashlib.sha256(template.read_bytes()).hexdigest()
        self.open_word()
        with patch.object(studio,'pick_import_file',return_value=template):self.page.locator('#word-pick-template').click();self.page.wait_for_function("document.querySelector('#word-template').value.includes('cover_template')")
        with patch.object(studio,'pick_project_parent',return_value=runtime['root']):
            self.page.locator('#word-pick-runtime').click();self.page.wait_for_function("document.querySelector('#word-export-status').textContent==='运行环境已更新。'")
        self.assertEqual(Path(self.page.locator('#word-template').input_value()).resolve(),template.resolve())
        output=self.compile()
        with zipfile.ZipFile(template) as old,zipfile.ZipFile(output) as new:
            for name in old.namelist():
                if name.startswith(('word/header','word/footer')):self.assertEqual(old.read(name),new.read(name))
            original=E.fromstring(old.read('word/document.xml'));result=E.fromstring(new.read('word/document.xml'))
            ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            slot=original.xpath('//w:body/w:p[w:bookmarkStart[@w:name="WPS_ReportBody"]]',namespaces=ns)[0]
            prefix=list(original.find('w:body',ns))[:slot.getparent().index(slot)]
            actual=list(result.find('w:body',ns))[:len(prefix)]
            canonical=lambda n:E.tostring(E.fromstring(E.tostring(n)),method='c14n')
            self.assertEqual([canonical(n) for n in prefix],[canonical(n) for n in actual])
            self.assertTrue(result.xpath('//w:instrText[contains(.,"TOC")]|//w:fldSimple[contains(@w:instr,"TOC")]',namespaces=ns))
        self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest(),before)
        self.open_word();self.assertEqual(Path(self.page.locator('#word-template').input_value()).resolve(),template.resolve())
        self.page.locator('#modal-cancel').click()
        count=len(list((self.project/'output').glob('*.docx')))
        # Real Bridge output policy rejects unsupported free HTML; no half DOCX.
        source(self.page,'# Advanced\n\n## Chapter\n\n<div style="display:flex">Card</div>\n')
        self.open_word();self.page.locator('#modal-ok').click()
        self.page.wait_for_function("document.querySelector('#word-export-status').textContent.startsWith('导出未完成')",timeout=90000)
        self.assertEqual(len(list((self.project/'output').glob('*.docx'))),count)
        self.assertTrue(self.page.locator('#modal-layer').is_visible())
        if self.evidence:self.page.screenshot(path=str(self.evidence/'unsupported-html.png'))
        self.page.locator('#modal-cancel').click()


    def test_04_managed_default_survives_original_move_and_new_project(self):
        source(self.page,TEXT);runtime=self.server.word_export.runtime()
        original=self.project.parent/'带封面的模板.docx';shutil.copyfile(runtime['wps']/'templates/cover_template_v0311.docx',original)
        expected=original.read_bytes();self.open_word()
        with patch.object(studio,'pick_import_file',return_value=original):
            self.page.locator('#word-pick-template').click();self.page.wait_for_function("document.querySelector('#word-template').value.includes('带封面')")
        self.page.locator('#word-set-default').click()
        self.page.wait_for_function("document.querySelector('#word-export-status').textContent.includes('默认模板已保管')",timeout=30000)
        managed=Path(self.page.locator('#word-template').input_value());self.assertEqual(managed.read_bytes(),expected)
        self.assertNotEqual(managed,original);original.rename(original.with_suffix('.moved.docx'))
        self.page.locator('#modal-cancel').click();self.open_word();self.assertEqual(Path(self.page.locator('#word-template').input_value()),managed)
        from word_export import WordExport
        restored=WordExport(studio.HERE,self.server.workspace)
        new=restored.describe('another-new-project');self.assertEqual(new['template_choice'],'managed');self.assertEqual(Path(new['template']),managed)
        self.assertEqual(restored.template('another-new-project','managed')[1],hashlib.sha256(expected).hexdigest())
        self.assertEqual(self.page.locator('#word-default-template').inner_text(),'使用程序原始模板')
        output=self.compile();self.assertTrue(output.is_file());self.assertEqual(managed.read_bytes(),expected)

    def test_05_references_abstract_and_strip_emphasis_are_explicit_output_policy(self):
        raw='# 输出策略\n\n> 摘要：摘要正文不应整体加粗。\n\n## 章节\n\n正文**加粗**和*中文斜体*，来源[^a]和重复[^a]。\n\n| 参数 | 结果 |\n| --- | --- |\n| 数据 | 12 |\n\n: 表题[^b] {#tbl-policy}\n\n[^a]: 第一条**参考文献**，[链接](https://example.com/)。\n\n[^b]: 第二条参考。\n'
        source(self.page,raw);self.open_word()
        self.assertFalse(self.page.locator('#word-keep-emphasis').is_checked())
        self.page.locator('#word-notes-mode').select_option('references')
        output=self.compile();doc=Document(output)
        heading=next(p for p in doc.paragraphs if p.text=='参考文献')
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        self.assertEqual(heading._p.xpath('./w:pPr/w:outlineLvl/@w:val'),['0'])
        self.assertEqual(heading._p.xpath('./w:pPr/w:numPr/w:numId/@w:val'),['0'])
        abstract=next(p for p in doc.paragraphs if p.text.startswith('摘要：'))
        self.assertTrue(abstract.runs[0].bold);self.assertTrue(all(r.bold is not True for r in abstract.runs[1:]))
        body=next(p for p in doc.paragraphs if p.text.startswith('正文'))
        self.assertTrue(all(r.bold is not True and r.italic is not True for r in body.runs))
        fields=doc._element.xpath('//w:fldSimple/@w:instr')
        self.assertEqual(sum('SEQ HRSReference' in f for f in fields),2);self.assertGreaterEqual(sum(' REF HR_' in f for f in fields),3)
        self.assertFalse(doc._element.xpath('//w:footnoteReference'))
        self.assertEqual(value(self.page),raw)
        meta=json.loads(Path(str(output)+'.hrs-bridge.json').read_text('utf8'))
        self.assertEqual(meta['word_output_options'],{'notes_mode':'references','preserve_emphasis':False})
        self.assertEqual(len(meta['derived_sections']),1)
        self.open_word();self.assertEqual(self.page.locator('#word-notes-mode').input_value(),'references');self.page.locator('#modal-cancel').click()

if __name__=='__main__':unittest.main()

"""WORD-LITE-01: same engines, embedded Python, open-page model and optional math."""
from pathlib import Path
import base64,hashlib,json,os,shutil,subprocess,unittest,zipfile
from unittest.mock import patch
from docx import Document
from PIL import Image
from browser_support import ManualService
from browser_support import source,value
import test_word_export_browser as previous
import studio


class WordLite(ManualService):
    open_word=previous.WordUI.open_word
    compile=previous.WordUI.compile

    def setUp(self):
        super().setUp();self.docid='f'*32;self.server.registry.register(self.docid,self.project)
        self.page.evaluate('id=>ReportStudioProgrammatic.open(id)',self.docid);self.page.set_default_timeout(20000)

    def failed(self):
        self.page.wait_for_function("document.querySelector('#word-export-status')?.textContent.startsWith('导出未完成')",timeout=90000)
        self.assertFalse(list((self.project/'output').glob('*.docx')))
        self.assertFalse(self.server.word_export.jobs)
        return self.page.locator('#word-export-status').inner_text()

    def test_01_formula_is_optional_then_real_omml(self):
        source(self.page,'# 原生公式\n\n## 计算\n\n向量 $\\boldsymbol{E}=mc^2$。\n\n$$\\frac{a}{b}+x^2$$\n\n```tex\n$not_math$\n```\n')
        pandoc=Path(os.environ['HRS_WORD_TEST_ENV'])/'pandoc/pandoc.exe';self.assertTrue(pandoc.is_file())
        exporter=self.server.word_export
        # Hide the Full test runner's formula tool until the user selects it.
        def selected(rt):return Path(exporter.settings['pandoc']) if exporter.settings.get('pandoc') else None
        with patch.object(exporter,'pandoc',side_effect=selected):
            self.open_word();self.page.locator('#modal-ok').click()
            self.assertIn('选择 Pandoc',self.failed())
            self.page.locator('.word-export-formula>summary').click()
            with patch.object(studio,'pick_import_file',return_value=pandoc):
                self.page.locator('#word-pick-formula').click()
                self.page.wait_for_function("document.querySelector('#word-export-status').textContent==='原生公式转换器已配置。'")
            output=self.compile()
        doc=Document(output);ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        self.assertEqual(len(doc._element.xpath('//m:oMath')),2)
        self.assertTrue(any('$not_math$' in p.text for p in doc.paragraphs))

    def rich_source(self):
        Image.new('RGB',(60,30),(40,120,190)).save(self.project/'assets/a.png')
        Image.new('RGB',(45,40),(170,75,35)).save(self.project/'assets/b.png')
        svg='<svg xmlns="http://www.w3.org/2000/svg" width="240" height="90" viewBox="0 0 240 90"><rect width="240" height="90" fill="#3461af"/><text x="20" y="55" fill="#fff" font-size="24">SVG</text></svg>'
        (self.project/'assets/vector.svg').write_bytes(svg.encode('utf8'))
        group={'schema':'report-image-group-v1','id':'fig-pair','caption':'两张图','columns':2,'rows':1,
               'images':[{'src':'assets/a.png','caption':'甲'},{'src':'assets/b.png','caption':'乙'}]}
        group_source=self.page.evaluate('d=>ReportLayout.groupSnippet(d)',group)
        figures=[{'schema':'report-figure-v1','id':'fig-flow','caption':'多行流程图','kind':'mermaid','source':'flowchart TB\n A["输入<br/>第二行"] --> B["输出"]'},
                 {'schema':'report-figure-v1','id':'fig-vector','caption':'矢量图','kind':'image','src':'assets/vector.svg','alt':'SVG'}]
        table={'schema':'report-table-v1','id':'tbl-params','caption':'参数','rows':[[{'text':'参数'},{'text':'数值'}],[{'text':'功率'},{'text':'12'}]]}
        for row in table['rows']:
            for cell in row:cell.update(rowspan=1,colspan=1)
        return '# 精简组件报告\n\n> 摘要：沿用同一工程编译规则。\n\n## 方案\n\n==重点==，`代码`与普通正文。\n\n1. 甲\n2. 乙\n\n'+group_source+'\n\n'+''.join('```report-figure\n'+json.dumps(f,ensure_ascii=False)+'\n```\n\n' for f in figures)+'```report-table\n'+json.dumps(table,ensure_ascii=False)+'\n```\n\n参见 @fig:fig-flow、@fig:fig-pair、@tbl:tbl-params。\n'

    def test_02_graphics_without_extra_browser_and_existing_cli_parity(self):
        text=self.rich_source();source(self.page,text)
        # A GUI export must succeed even if Node/second-browser lookup is broken.
        # These vars are changed only after the test's own visible-page browser starts.
        with patch.dict(os.environ,{'HRS_BRIDGE_NODE':'Z:/no-node.exe','CHROMIUM_PATH':'Z:/no-browser.exe'}),patch.object(self.server.word_export,'pandoc',return_value=None):
            self.open_word();output=self.compile()
        doc=Document(output);self.assertEqual(len(doc.inline_shapes),4);self.assertEqual(len(doc.tables),2)
        instructions=doc._element.xpath('//w:instrText/text()|//w:fldSimple/@w:instr')
        self.assertTrue(any('SEQ' in x for x in instructions));self.assertTrue(any('REF' in x for x in instructions))
        self.assertEqual((self.project/'report.md').read_bytes(),b'# Initial\n')
        rt=self.server.word_export.runtime();env=Path(os.environ['HRS_WORD_TEST_ENV'])
        cli_source=self.project.parent/'cli-source';shutil.copytree(self.project/'assets',cli_source/'assets')
        (cli_source/'report.md').write_bytes(text.encode('utf8'))
        (cli_source/'meta.json').write_text(json.dumps({'render_profile':'free','render_options':{'show_object_numbers':True}}),encoding='utf8')
        cli_out=self.project.parent/'cli.docx';app=Path(studio.__file__).parent
        bootstrap='import runpy,sys;sys.path.insert(0,'+repr(str(rt['bridge']))+');runpy.run_module("hrs_bridge",run_name="__main__")'
        command=[str(env/'python-bridge/python.exe'),'-X','utf8','-B','-c',bootstrap,'--hrs-root',str(app),'--wps-root',str(rt['wps']),
                 '--node',str(env/'node/node.exe'),'--browser',str(env/'browser/chrome-win/chrome.exe'),'to-docx',str(cli_source),'--out',str(cli_out)]
        done=subprocess.run(command,capture_output=True,text=True,encoding='utf8',timeout=90)
        self.assertEqual(done.returncode,0,done.stderr)
        cli=Document(cli_out)
        self.assertEqual([p.text for p in doc.paragraphs],[p.text for p in cli.paragraphs])
        self.assertEqual(len(cli.inline_shapes),4);self.assertEqual(len(cli.tables),len(doc.tables))
        self.assertEqual(instructions,cli._element.xpath('//w:instrText/text()|//w:fldSimple/@w:instr'))
        for file in (output,cli_out):
            meta=json.loads(Path(str(file)+'.hrs-bridge.json').read_text('utf8'))
            self.assertEqual(base64.b64decode(meta['source_b64']).decode('utf8'),text)
        if self.evidence:
            shutil.copyfile(cli_out,self.evidence/'existing-cli.docx')
            self.page.screenshot(path=str(self.evidence/'graphics-preview.png'))

    def test_03_model_and_asset_changes_reject_without_partial_output(self):
        text=self.rich_source();source(self.page,text);self.open_word()
        def tamper(route):
            req=route.request.post_data_json;req['model']['request_sha256']='0'*64
            route.continue_(post_data=json.dumps(req))
        self.page.route('**/api/word/export',tamper);self.page.locator('#modal-ok').click()
        self.assertIn('快照不一致',self.failed());self.page.unroute('**/api/word/export',tamper)
        def change_asset(route):
            Image.new('RGB',(60,30),(9,9,9)).save(self.project/'assets/a.png');route.continue_()
        self.page.route('**/api/word/export',change_asset);self.page.locator('#modal-ok').click()
        self.assertIn('已改变',self.failed());self.page.unroute('**/api/word/export',change_asset)
        self.assertEqual(value(self.page),text);self.assertEqual((self.project/'report.md').read_bytes(),b'# Initial\n')
        self.page.locator('#modal-cancel').click()


if __name__=='__main__':unittest.main()

"""Map the current HRS IR into the independently configured WPS template."""
import base64
from copy import deepcopy
from io import BytesIO
import html
import json
from pathlib import Path
import tempfile

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from docx.shared import Twips, Emu
from docx.text.paragraph import Paragraph
from lxml import etree as E
from PIL import Image

from .browser import call
from .math_adapter import MathAdapter
from .ooxml import N, W, element, zero_indent, field, semantic_hash, provenance_path, drawing_extents
from .runtime import BridgeError, new_path, sha, source_contract_valid

def output_options(value=None):
    """Output policy is snapshot-bound and never rewrites Markdown."""
    value={} if value is None else value
    if not isinstance(value,dict) or set(value)-{'notes_mode','preserve_emphasis'}:raise BridgeError('Word 输出选项无效')
    mode=value.get('notes_mode','footnotes');keep=value.get('preserve_emphasis',True)
    if mode not in ('footnotes','references') or type(keep) is not bool:raise BridgeError('Word 脚注或正文强调选项无效')
    return {'notes_mode':mode,'preserve_emphasis':keep}


class NoteBody:
    """Use python-docx paragraphs with the note's own OPC relationships."""
    def __init__(self, node, part):
        self._element=node
        self._part = part

    @property
    def part(self):
        return self._part

    def add_paragraph(self):
        node=element('w:p');self._element.append(node)
        return Paragraph(node,self)


class Writer:
    def __init__(self, rt, source, template=None, accept_numbered_headings=False, *, model=None, prepare_only=False, options=None):
        from wps_pipeline.resolve_styles import inspect_template, open_template_body
        self.rt = rt
        self.options=output_options(options)
        self.derived_sections=[]
        self.source = Path(source).resolve()
        if self.source.is_file():
            self.source = self.source.parent
        self.raw = (self.source / 'report.md').read_bytes()
        from render_profile import read as read_profile, read_options, require_output
        self.source_profile=read_profile(self.source);self.source_options=read_options(self.source)
        self.document_key = sha(self.raw)[:16]
        self.ir = rt.analyze(self.raw, self.source / 'assets')
        if not self.ir['ok'] or not source_contract_valid(self.ir):
            raise BridgeError('HRS validation prevents compilation: ' + json.dumps(self.ir['diagnostics'], ensure_ascii=False))
        self.output_diagnostics=require_output(self.ir,'engineering',accept_numbered_headings)
        config = json.loads((rt.wps / 'config.json').read_text('utf-8-sig'))
        self.template = Path(template).resolve() if template is not None else rt.wps / config.get('template', 'templates/target_template.docx')
        # The current WPS configuration resolves the actual mother template.
        if not self.template.is_file():
            raise BridgeError('Configured WPS template does not exist: ' + str(self.template))
        self.template_sha256=sha(self.template.read_bytes())
        self.template_info = inspect_template(self.template)
        self.styles = self.template_info['styles']
        self.width = self.template_info['usable_width_twips']
        self.doc = Document(self.template)
        self.body = self.doc._element.body
        self.title_pattern=open_template_body(self.body,self.template_info)
        self.front_nodes=list(self.body)[:-1]
        self.front_xml=[E.tostring(node) for node in self.front_nodes]
        self.bind_heading_numbering()
        self.heading_counts = [0] * 6
        self.math = None
        self.note_part = None
        self.note_refs = {}
        self.current_parent = None
        self.bookmark_id = max([int(n.get(qn('w:id'))) for n in self.body.iter(qn('w:bookmarkStart')) if (n.get(qn('w:id')) or '').isdigit()]+[0])
        self.tables = 0
        self.diagnostics = list(self.output_diagnostics)
        self.emf_parts = {}
        self.inline_text = []
        self.block_text = []
        self.block_starts = []
        self.inline_index = {}
        self.markdown_parts = {}
        self.metadata_blocks = []
        self.origin_styles = {}
        self.word_origin = None
        self.image_layouts = {}
        origin_path = self.source / 'bridge-origin.json'
        if origin_path.is_file():
            origin = json.loads(origin_path.read_text('utf8'))
            if origin.get('schema') == 'hrs-word-origin-v1' and origin.get('markdown_sha256') == sha(self.raw):
                self.word_origin = origin
                self.origin_styles = {b['start']:b for b in origin.get('word_styles',[])}
                self.image_layouts = {b['id']:b for b in origin.get('image_layouts',[])}
        if self.word_origin is None:
            headings = [b for b in self.ir['blocks'] if b['type'] == 'heading']
            titles = [b for b in headings if b['level'] == 1]
            if len(titles) > 1 or (titles and headings[0] is not titles[0]):
                raise BridgeError('AMBIGUOUS_SINGLE_REPORT_TITLE: a single-report H1 title must precede its chapters; organize multiple reports explicitly')
        self.heading_output = {}
        for index, b in enumerate(self.ir['blocks']):
            b['index'] = index  # snapshot-local bridge receipt, never a persistent source ID
            if b['type'] == 'markdown':
                self.markdown_parts[index] = self.project_markdown(b)
            texts = list(self.texts(b))
            if b['type'] == 'heading':
                style = self.heading_style(b)
                text = b['text'] if style == 'Title' else self.heading_text(b, int(style.split()[-1]))
                self.heading_output[index] = (style, text)
                texts.append(text)
            for text in texts:
                if text not in self.inline_index:
                    self.inline_index[text] = len(self.inline_text)
                    self.inline_text.append(text)
        self.model_request={'schema':'hrs-word-model-request-v1','markdown':self.raw.decode('utf8'),
                            'sections':self.block_text,'sectionStarts':self.block_starts,'inline':self.inline_text,
                            'template_sha256':self.template_sha256,'source_sha256':sha(self.raw),'output_options':self.options,
                            'assets':{name:sha((self.source/name).read_bytes()) for name in sorted(self.ir['assets'])},
                            'diagram_sources':{b['id']:b['data']['source'] for b in self.ir['blocks'] if b['type']=='figure' and b['data']['kind']=='mermaid'},
                            'svg_sources':{Path(name).name:base64.b64encode((self.source/name).read_bytes()).decode('ascii') for name in sorted(self.ir['assets']) if name.lower().endswith('.svg')}}
        self.model_request['request_sha256']=sha(json.dumps(self.model_request,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode('utf8'))
        self.model=None
        if not prepare_only:
            self.model=model if model is not None else call(rt,'model',project=self.source,**self.model_request)
            self.validate_model()
            self.diagnostics.extend(self.model['profile']['diagnostics'])
        self.identities = {b['id']: b for b in self.ir['blocks'] if b.get('id')}

    def validate_model(self):
        """The editor model is bound to this exact source/assets/template request."""
        m=self.model
        if not isinstance(m,dict) or m.get('schema')!='hrs-word-browser-model-v1' or m.get('request_sha256')!=self.model_request['request_sha256']:
            raise BridgeError('Word 页面模型与当前源稿、附件或模板快照不一致，请重新导出')
        if m.get('output_options')!=self.options:raise BridgeError('Word 页面模型与当前输出策略不一致，请重新导出')
        for field,expected in (('tokens',len(self.block_text)),('inline',len(self.inline_text))):
            values=m.get(field)
            if not isinstance(values,list) or len(values)!=expected or any(not isinstance(v,list) for v in values):
                raise BridgeError('Word 页面模型的段落范围不完整')
        if not isinstance(m.get('profile'),dict) or not isinstance(m['profile'].get('diagnostics'),list):
            raise BridgeError('Word 页面模型缺少工程预检信息')
        notes=m.get('footnotes',{})
        if not isinstance(notes,dict) or len(notes)>1000 or any(not isinstance(k,str) or not isinstance(v,list) for k,v in notes.items()):
            raise BridgeError('Word 脚注模型无效')
        def check_notes(value,in_note=False):
            if isinstance(value,list):
                for item in value:check_notes(item,in_note)
            elif isinstance(value,dict):
                if value.get('type')=='hrsFootnoteRef':
                    if in_note:raise BridgeError('Word 不支持脚注中嵌套脚注')
                    if value.get('key') not in notes:raise BridgeError('脚注未定义：'+str(value.get('key')))
                if in_note and value.get('type')=='image':raise BridgeError('Word 脚注暂不支持图片')
                for key,item in value.items():
                    if key not in ('raw','text'):check_notes(item,in_note)
        check_notes(m['tokens']);check_notes(m['inline'])
        for value in notes.values():check_notes(value,True)
        for field,expected in (('diagrams',self.model_request['diagram_sources']),('svgAssets',self.model_request['svg_sources'])):
            values=m.get(field)
            if not isinstance(values,dict) or not set(expected)<=set(values):raise BridgeError('Word 图形预览不完整')
            for ident,encoded in values.items():
                if not isinstance(encoded,str) or len(encoded)>16*1024*1024:raise BridgeError('Word 图形预览大小无效')
                try:
                    raw=base64.b64decode(encoded,validate=True)
                    with Image.open(BytesIO(raw)) as image:
                        if image.format!='PNG' or image.width*image.height>16000000:raise ValueError()
                        image.verify()
                except Exception:raise BridgeError('Word 图形预览不是有效的有界 PNG') from None
        def math_node(value):
            if isinstance(value,dict):return value.get('type')=='hrsMath' or any(math_node(v) for k,v in value.items() if k not in ('raw','text'))
            return isinstance(value,list) and any(math_node(v) for v in value)
        if not self.rt.pandoc and (any(b['type']=='equation' for b in self.ir['blocks']) or math_node(m['tokens']) or math_node(m['inline']) or math_node(notes)):
            raise BridgeError('本文含公式；请在“公式支持”中选择 Pandoc 后重试。普通无公式报告无需该组件。')

    @staticmethod
    def texts(block):
        if block['type'] == 'heading':
            yield block['text']
        data = block.get('data', {})
        if 'caption' in data:
            yield data['caption']
        for item in data.get('images', []):
            yield item.get('caption', '')
        for row in data.get('rows', []) if isinstance(data.get('rows'), list) else []:
            for cell in row:
                if cell is not None:
                    yield cell['text'] if isinstance(cell, dict) else cell

    def diagnostic(self, code, block, message):
        self.diagnostics.append({'code': code, 'block': block['index'], 'id': block.get('id'), 'message': message})

    def paragraph(self, style='Normal', parent=None, compact=False):
        from wps_pipeline.docx_formatting import style_paragraph
        parent=parent or self.current_parent
        if isinstance(parent,NoteBody) and style=='Normal':style='Footnote Text'
        p = self.doc.add_paragraph() if parent is None else parent.add_paragraph()
        # WPS helpers use plain lxml, python-docx uses registered OOXML subclasses.
        # Cross that boundary explicitly; do not fork the shared formatting rules.
        plain = E.fromstring(E.tostring(p._p))
        if style.startswith('HRS ') and style not in self.doc.styles:
            raise BridgeError('Current Word template requires '+style+'; prepare template v0.3.11 before export')
        style_id=self.styles.get(style,self.doc.styles[style].style_id if style in self.doc.styles else style)
        style_paragraph(plain, style_id, compact=compact)
        formatted = parse_xml(E.tostring(plain))
        p._p.getparent().replace(p._p, formatted)
        p = Paragraph(formatted, p._parent)
        if style in ('图形段', '图题', '表题', '表内正文', '公式'):
            zero_indent(p._p)
        return p

    def inline(self, p, tokens, bold=False, italic=False, strike=False, pre_wrap=False):
        from report_contract import REF
        from reference_context import masked_context
        for t in tokens:
            kind = t['type']
            if kind in ('strong', 'em', 'del'):
                self.inline(p, t['tokens'], bold or kind == 'strong' and self.options['preserve_emphasis'], italic or kind == 'em' and self.options['preserve_emphasis'], strike or kind == 'del', pre_wrap)
            elif kind in ('text', 'escape'):
                if t.get('tokens'):
                    self.inline(p, t['tokens'], bold, italic, strike, pre_wrap)
                    continue
                text = t.get('text', '')
                cursor = 0
                for ref in REF.finditer(masked_context(text)):
                    self.run(p, text[cursor:ref.start()], bold, italic, strike, pre_wrap)
                    target = self.identities.get(ref[2])
                    if target is None:
                        raise BridgeError('Reference target missing: ' + ref[2])
                    if ref[1] in ('fig', 'tbl'):
                        self.run(p, '图 ' if ref[1] == 'fig' else '表 ', bold, italic, strike)
                        field(p, ' REF ' + self.number_name(ref[2]) + ' \\h ', self.output_numbers[ref[2]])
                    elif ref[1] == 'sec':
                        field(p, ' REF ' + self.number_name(ref[2]) + ' \\h ', target['text'])
                    else:
                        raise BridgeError('WPS profile has unnumbered equations; equation references require an explicit output policy')
                    cursor = ref.end()
                self.run(p, text[cursor:], bold, italic, strike, pre_wrap)
            elif kind == 'codespan':
                self.run(p, t['text'], bold, italic, strike)
            elif kind == 'hrsMath':
                if self.math is None:
                    self.math = MathAdapter(self.rt)
                p._p.append(self.math.omml(t['tex'], False))
            elif kind == 'hrsFootnoteRef':
                self.footnote(p,t['key'])
            elif kind == 'br':
                p.add_run().add_break()
            elif kind == 'link':
                from docx.opc.constants import RELATIONSHIP_TYPE as RT
                hyperlink = element('w:hyperlink')
                hyperlink.set(qn('r:id'), p.part.relate_to(t['href'], RT.HYPERLINK, is_external=True))
                start = len(p._p)
                self.inline(p, t.get('tokens', [{'type': 'text', 'text': t['text']}]), bold, italic, strike, pre_wrap)
                for node in list(p._p)[start:]:
                    hyperlink.append(node)
                p._p.append(hyperlink)
            elif kind == 'image':
                self.picture(p, {'src': t['href'], 'alt': t.get('text', '')}, self.width)
                zero_indent(p._p)
            else:
                raise BridgeError('Unsupported inline Markdown token: ' + kind)

    def footnote(self,p,key):
        """One native footnote per definition; later mentions are live NOTEREF."""
        from docx.opc.constants import RELATIONSHIP_TYPE as RT, CONTENT_TYPE as CT
        from docx.opc.part import XmlPart
        from docx.opc.packuri import PackURI
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt
        if self.current_parent is not None:raise BridgeError('Word 不支持脚注中嵌套脚注')
        if self.options['notes_mode']=='references':
            if key not in self.note_refs:self.note_refs[key]=('HR_'+self.document_key+'_'+sha(key.encode('utf8'))[:16],len(self.note_refs)+1)
            name,number=self.note_refs[key]
            self.run(p,'[');field(p,' REF '+name+' \\h ',number);self.run(p,']')
            return
        if key in self.note_refs:
            name,number=self.note_refs[key]
            self.note_bracket(p,'[')
            field(p,' NOTEREF '+name+' \\h \\f ',number)
            run=p._p[-1].find(qn('w:r'))
            props=element('w:rPr');props.append(element('w:rStyle',val=self.doc.styles['Footnote Reference'].style_id));props.append(element('w:vertAlign',val='superscript'));run.insert(0,props)
            self.note_bracket(p,']')
            return
        if self.note_part is None:
            for name,kind in [('Footnote Text',WD_STYLE_TYPE.PARAGRAPH),('Footnote Reference',WD_STYLE_TYPE.CHARACTER)]:
                if name in self.doc.styles:continue
                style=self.doc.styles.add_style(name,kind)
                if kind==WD_STYLE_TYPE.PARAGRAPH:
                    style.base_style=self.doc.styles['Normal'];style.font.size=Pt(10)
                    style.paragraph_format.first_line_indent=Twips(0)
                    style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(0)
                    style.paragraph_format.line_spacing=1
                else:style.font.superscript=True
            try:
                old=self.doc.part.part_related_by(RT.FOOTNOTES)
                self.note_part=XmlPart(old.partname,old.content_type,parse_xml(old.blob),old.package)
                for rid,rel in old.rels.items():
                    self.note_part.rels.add_relationship(rel.reltype,rel.target_ref if rel.is_external else rel.target_part,rid,rel.is_external)
                for rel in list(self.doc.part.rels.values()):
                    if rel.reltype==RT.FOOTNOTES:self.doc.part.rels.add_relationship(rel.reltype,self.note_part,rel.rId)
            except KeyError:
                root=parse_xml(('<w:footnotes xmlns:w="'+N['w']+'" xmlns:r="'+N['r']+'"/>').encode())
                self.note_part=XmlPart(PackURI('/word/footnotes.xml'),CT.WML_FOOTNOTES,root,self.doc.part.package)
                self.doc.part.relate_to(self.note_part,RT.FOOTNOTES)
            root=self.note_part.element
            for kind,ident in [('separator',-1),('continuationSeparator',0)]:
                if not any(n.get(qn('w:type'))==kind for n in root):
                    note=element('w:footnote',type=kind,id=ident);para=element('w:p');run=element('w:r');run.append(element('w:'+kind));para.append(run);note.append(para);root.insert(0,note)
        root=self.note_part.element
        ident=max([int(n.get(qn('w:id'))) for n in root if (n.get(qn('w:id')) or '').lstrip('-').isdigit()]+[0])+1
        number=len(self.doc._element.xpath('//w:footnoteReference'))+1
        name='HF_'+self.document_key+'_'+sha(key.encode('utf8'))[:16]
        self.note_refs[key]=(name,number)
        run=element('w:r');props=element('w:rPr');props.append(element('w:rStyle',val=self.doc.styles['Footnote Reference'].style_id));run.append(props)
        run.append(element('w:footnoteReference',id=ident))
        self.note_bracket(p,'[');self.mark_inline(p,name,lambda:p._p.append(run));self.note_bracket(p,']')
        note=element('w:footnote',id=ident);root.append(note)
        self.current_parent=NoteBody(note,self.note_part)
        try:self.markdown(self.model['footnotes'][key])
        finally:self.current_parent=None
        paragraphs=note.findall(qn('w:p'))
        if not paragraphs:
            empty=element('w:p');note.append(empty);paragraphs=[empty]
        first=paragraphs[0];marker=element('w:r');props=element('w:rPr');props.append(element('w:rStyle',val=self.doc.styles['Footnote Reference'].style_id));marker.append(props)
        marker.append(element('w:footnoteRef'))
        space=element('w:r');text=element('w:t');text.set('{http://www.w3.org/XML/1998/namespace}space','preserve');text.text=' ';space.append(text)
        at=1 if first.find(qn('w:pPr')) is not None else 0
        first.insert(at,marker);first.insert(at+1,space)

    def note_bracket(self,p,text):
        run=p.add_run(text);run.style=self.doc.styles['Footnote Reference']

    def write_references(self):
        if self.options['notes_mode']!='references' or not self.note_refs:return
        self.bookmark_id+=1;ident=self.bookmark_id;name='HB_'+self.document_key+'_references'
        begin=element('w:bookmarkStart',id=ident,name=name);self.body.insert(len(self.body)-1,begin)
        heading=self.paragraph('Heading 1');heading.add_run('参考文献')
        # The existing TOC consumes its real outline level; numId=0 keeps this
        # unnumbered heading out of the template's chapter sequence.
        pr=heading._p.get_or_add_pPr()
        for key in ('numPr','outlineLvl'):
            for old in list(pr.findall(qn('w:'+key))):pr.remove(old)
        number=element('w:numPr');number.append(element('w:numId',val=0));pr.append(number);pr.append(element('w:outlineLvl',val=0))
        for key,(target,index) in self.note_refs.items():
            before=len(self.body);self.markdown(self.model['footnotes'][key])
            nodes=list(self.body)[before-1:-1];paragraphs=[n for n in nodes if n.tag==qn('w:p')]
            if not paragraphs:raise BridgeError('参考文献条目没有可输出的正文')
            first=Paragraph(paragraphs[0],self.doc._body)
            marker=self.paragraph();self.run(marker,'[')
            self.mark_inline(marker,target,lambda:field(marker,' SEQ HRSReference \\* ARABIC ',index));self.run(marker,'] ')
            at=1 if first._p.find(qn('w:pPr')) is not None else 0
            for node in list(marker._p):
                if node.tag!=qn('w:pPr'):first._p.insert(at,node);at+=1
            marker._p.getparent().remove(marker._p)
            zero_indent(first._p)
        end=element('w:bookmarkEnd',id=ident);self.body.insert(len(self.body)-1,end)
        nodes=list(self.body)[list(self.body).index(begin)+1:list(self.body).index(end)]
        self.derived_sections.append({'kind':'references','bookmark':name,'native_sha256':semantic_hash(nodes,self.doc.part.related_parts,self.doc.part.rels)})

    @staticmethod
    def run(p, text, bold=False, italic=False, strike=False, pre_wrap=False):
        if text:
            # Marked's soft source newline is whitespace; only its explicit br
            # token is a native line break. This also prevents a justified REF
            # line from stretching its number to the right margin in WPS.
            text=html.unescape(text)
            run = p.add_run(text.replace('\r\n','\n').replace('\r','\n') if pre_wrap else text.replace('\r\n',' ').replace('\n',' '))
            if bold: run.bold = True
            if italic: run.italic = True
            if strike: run.font.strike = True

    def rich(self, p, text, pre_wrap=False):
        self.inline(p, self.model['inline'][self.inline_index[text]], pre_wrap=pre_wrap)

    def markdown(self, tokens, level=0, list_style=None, quote_lines=False):
        for token in tokens:
            kind = token['type']
            if kind == 'space':
                continue
            if kind in ('paragraph', 'text'):
                p = self.paragraph()
                if quote_lines:
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    zero_indent(p._p)
                self.inline(p, token.get('tokens', [{'type': 'text', 'text': token.get('text', '')}]),pre_wrap=quote_lines)
            elif kind == 'blockquote':
                self.markdown(token['tokens'], level,quote_lines=True)
            elif kind == 'hrsAbstract':
                for part in token['tokens']:
                    if part['type']=='space':continue
                    p=self.paragraph()
                    self.inline(p,part.get('tokens',[{'type':'text','text':part.get('text','')}]))
                    # The label alone is bold. Body keeps the template's normal
                    # paragraph/font; HRS Abstract inherits chapter boldness.
                    import re
                    for run in p.runs:
                        match=re.match(r'^(\s*摘要[:：])',run.text)
                        if match:
                            from docx.text.run import Run
                            suffix=run.text[match.end():]
                            if suffix:
                                rest=Run(deepcopy(run._r),p);rest.text=suffix;run._r.addnext(rest._r)
                            run.text=match[1];run.bold=True
                        if run.text:break
            elif kind == 'code':
                p = self.paragraph('HRS Code')
                p.add_run(token['text'])
                zero_indent(p._p)
            elif kind == 'hr':
                p = self.paragraph()
                pr = p._p.get_or_add_pPr()
                borders = element('w:pBdr')
                borders.append(element('w:bottom', val='single', sz=4, color='auto'))
                pr.append(borders)
            elif kind == 'hrsStyledList':
                self.markdown(token['tokens'], level, token['style'])
            elif kind == 'list':
                num = self.list_numbering(token.get('ordered', False), token.get('start') or 1, list_style)
                for item in token['items']:
                    if item.get('task'):
                        raise BridgeError('Task checkbox Markdown requires an explicit native output mapping')
                    for child in item['tokens']:
                        if child['type'] == 'space': continue
                        if child['type'] == 'list':
                            self.markdown([child], level + 1)
                            continue
                        if child['type'] not in ('text', 'paragraph'):
                            raise BridgeError('Unsupported block inside a list: ' + child['type'])
                        p = self.paragraph()
                        self.inline(p, child.get('tokens', [{'type': 'text', 'text': child.get('text', '')}]))
                        pr = p._p.get_or_add_pPr()
                        number = element('w:numPr')
                        number.extend([element('w:ilvl', val=min(level, 8)), element('w:numId', val=num)])
                        pr.append(number)
                        zero_indent(p._p)
                        indent = pr.find(qn('w:ind'))
                        indent.set(qn('w:left'), str((level + 1) * 420))
                        indent.set(qn('w:hanging'), '210')
            else:
                # Raw HTML and unknown extensions must not silently disappear.
                raise BridgeError('Unsupported ordinary Markdown block: ' + kind)

    def list_numbering(self, ordered, start, style=None):
        # Project the five existing HRS display styles into Word numbering.
        # The original Marked extension owns the source syntax and list parsing.
        styles = {'bullet':('bullet','•'), 'decimal':('decimal','%n.'),
                  'lower-alpha':('lowerLetter','%n.'), 'paren-decimal':('decimal','（%n）'),
                  'paren-lower-alpha':('lowerLetter','（%n）')}
        if style is not None and style not in styles: raise BridgeError('Unsupported HRS list display style')
        fmt, label = styles[style] if style else (('decimal','%n.') if ordered else ('bullet','•'))
        root = self.doc.part.numbering_part.element
        abstract = max([int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))] + [-1]) + 1
        number = max([int(e.get(qn('w:numId'))) for e in root.findall(qn('w:num'))] + [0]) + 1
        a = element('w:abstractNum', abstractNumId=abstract)
        a.append(element('w:multiLevelType', val='multilevel'))
        for level in range(9):
            lvl = element('w:lvl', ilvl=level)
            lvl.extend([element('w:start', val=start), element('w:numFmt', val=fmt),
                        element('w:lvlText', val=label.replace('%n','%'+str(level+1)))])
            a.append(lvl)
        root.append(a)
        n = element('w:num', numId=number)
        n.append(element('w:abstractNumId', val=abstract))
        root.append(n)
        return number

    def bind_heading_numbering(self):
        """Word output adapter for the existing 1 / 1.1 / ... template policy.

        The mother template deliberately has no live heading list. Bind only
        the six output Heading styles; keep its other style properties intact.
        """
        root = self.doc.part.numbering_part.element
        # v0.3.11 owns the live outline. Reuse it without rebuilding definitions.
        bindings=[self.doc.styles['Heading '+str(i)].element.find('.//'+qn('w:numId')) for i in range(1,7)]
        if all(n is not None for n in bindings) and len({n.get(qn('w:val')) for n in bindings})==1:
            number=root.find(qn('w:num')+'[@'+qn('w:numId')+'="'+bindings[0].get(qn('w:val'))+'"]')
            aid=number.find(qn('w:abstractNumId')) if number is not None else None
            abstract=root.find(qn('w:abstractNum')+'[@'+qn('w:abstractNumId')+'="'+aid.get(qn('w:val'))+'"]') if aid is not None else None
            levels=abstract.findall(qn('w:lvl')) if abstract is not None else []
            if len(levels)==6 and all(level.find(qn('w:lvlText')) is not None and level.find(qn('w:lvlText')).get(qn('w:val'))=='.'.join('%'+str(j) for j in range(1,i+2)) for i,level in enumerate(levels)):
                return
        aid = max([int(e.get(qn('w:abstractNumId'))) for e in root.findall(qn('w:abstractNum'))] + [-1]) + 1
        nid = max([int(e.get(qn('w:numId'))) for e in root.findall(qn('w:num'))] + [0]) + 1
        abstract = element('w:abstractNum', abstractNumId=aid)
        abstract.append(element('w:multiLevelType', val='multilevel'))
        abstract.append(element('w:name', val='HRS Template Headings'))
        for level in range(6):
            style = self.doc.styles['Heading ' + str(level + 1)]
            lvl = element('w:lvl', ilvl=level)
            lvl.extend([element('w:start', val=1), element('w:numFmt', val='decimal')])
            if level:
                lvl.append(element('w:lvlRestart', val=level))
            lvl.extend([element('w:pStyle', val=style.style_id), element('w:suff', val='space'),
                        element('w:lvlText', val='.'.join('%' + str(i) for i in range(1, level + 2))),
                        element('w:lvlJc', val='left')])
            p = element('w:p')
            zero_indent(p)
            lvl.append(p.get_or_add_pPr())
            rp = element('w:rPr')
            rp.append(element('w:rFonts', ascii='Times New Roman', hAnsi='Times New Roman'))
            lvl.append(rp)
            abstract.append(lvl)
            np = style.element.get_or_add_pPr().get_or_add_numPr()
            np.get_or_add_ilvl().val = level
            np.get_or_add_numId().val = nid
        root.insert(len(root.findall(qn('w:abstractNum'))), abstract)
        number = element('w:num', numId=nid)
        number.append(element('w:abstractNumId', val=aid))
        root.append(number)

    def heading_style(self, block):
        origin = self.origin_styles.get(block['start'], {})
        if origin.get('source_sha256') == sha(block['source'].encode('utf8')):
            style=origin.get('style')
            if style=='Title' or style in ['Heading '+str(i) for i in range(1,7)]:return style
        if self.word_origin and any(v.get('projection') for v in self.origin_styles.values()):
            return 'Title' if block['level']==1 else 'Heading '+str(block['level']-1)
        if self.word_origin is not None:
            return 'Heading ' + str(block['level'])
        # HRS single report: H1 is its title; H2 starts the first body level.
        return 'Title' if block['level'] == 1 else 'Heading ' + str(block['level'] - 1)

    def project_markdown(self, block):
        """Split only verified derived ranges; all other content uses the original HRS lexer."""
        raw=block['source'].encode('utf-16-le');cursor=0;parts=[]
        def markdown(data):
            if data:
                parts.append(('markdown',len(self.block_text)))
                self.block_text.append(data.decode('utf-16-le'))
                self.block_starts.append(block['start']+cursor//2)
        for start,hint in sorted(self.origin_styles.items()):
            if hint.get('projection')!='heading-list' or not block['start']<=start<block['end']:continue
            lo=(start-block['start'])*2;hi=(hint['end']-block['start'])*2
            source=raw[lo:hi].decode('utf-16-le')
            if lo<cursor or hi>len(raw) or sha(source.encode('utf8'))!=hint['source_sha256'] or not source.startswith('- '):
                raise BridgeError('Invalid source-bound heading projection')
            markdown(raw[cursor:lo]);text=source[2:].rstrip('\r\n')
            style=hint['style']
            if style!='Heading 6':raise BridgeError('Unsupported overflow heading role')
            text=self.heading_text(dict(block,text=text),6)
            if text not in self.inline_index:
                self.inline_index[text]=len(self.inline_text);self.inline_text.append(text)
            parts.append((style,text));cursor=hi
        markdown(raw[cursor:]);return parts

    def heading_text(self, block, level):
        from wps_pipeline.numbered_candidates import numeric_expression,manual_prefix
        self.heading_counts[level - 1] += 1
        self.heading_counts[level:] = [0] * (6 - level)
        text = block['text']
        # Compare only an exact expected output label plus a literal space.
        # Reuse the WPS quantity/ambiguity guard; never infer engineering text.
        prefix = '.'.join(map(str, self.heading_counts[:level])) + ' '
        parsed=manual_prefix(text)
        if parsed and parsed['numbers']==self.heading_counts[:level]:prefix=parsed['prefix']
        if all(self.heading_counts[:level]) and text.startswith(prefix):
            if numeric_expression(text) is None:
                self.diagnostic('HEADING_LITERAL_PREFIX_TO_NATIVE', block,
                                'Exact chapter prefix moved to native Word outline numbering; source is retained.')
                return text[len(prefix):]
            self.diagnostic('HEADING_NUMERIC_TEXT_RETAINED', block,
                            'Numeric or ambiguous title text is retained; no chapter prefix deletion was inferred.')
        return text

    def picture(self, p, data, width, diagram=None):
        source = data.get('src')
        if diagram:
            blob = base64.b64decode(self.model['diagrams'][diagram])
        elif source and Path(source).suffix.lower() == '.svg':
            blob = base64.b64decode(self.model['svgAssets'][Path(source).name])
        else:
            path = (self.source / (source or '')).resolve()
            if not path.is_relative_to(self.source / 'assets'):
                raise BridgeError('Image must be inside project assets: ' + str(source))
            blob = path.read_bytes()
        with Image.open(BytesIO(blob)) as img:
            w, h = img.size
            # python-docx cannot embed WebP; this is a declared display rendition.
            if img.format not in ('PNG', 'JPEG', 'GIF', 'BMP', 'TIFF'):
                buffer = BytesIO()
                img.convert('RGBA').save(buffer, format='PNG')
                blob = buffer.getvalue()
        target_width = max(1, min(width, round(10000 * w / h)))
        shape = p.add_run().add_picture(BytesIO(blob), width=Twips(target_width))
        layout=self.image_layouts.get(data.get('id'))
        if layout and source and sha((self.source/source).read_bytes())==layout['source_sha256']:
            original_width,original_height=layout['width_emu'],layout['height_emu']
            if not isinstance(original_width,int) or not isinstance(original_height,int) or min(original_width,original_height)<=0:
                raise BridgeError('Invalid Word-origin image extent')
            scale=min(1,width*635/original_width,10000*635/original_height)
            shape.width=Emu(round(original_width*scale));shape.height=Emu(round(original_height*scale))
        shape._inline.docPr.set('descr', data.get('alt', data.get('caption', '')))
        if source:
            shape._inline.docPr.set('name', source)
        original = data.get('sourceAsset')
        if original and original.lower().endswith('.emf'):
            from docx.opc.packuri import PackURI
            from docx.opc.part import Part
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            emf = (self.source / original).read_bytes()
            digest = sha(emf)
            if digest not in self.emf_parts:
                self.emf_parts[digest] = Part(PackURI('/word/media/hrs_' + digest[:16] + '.emf'), 'image/x-emf', emf, self.doc.part.package)
            rid = self.doc.part.relate_to(self.emf_parts[digest], RT.IMAGE)
            shape._inline.xpath('.//a:blip')[0].set(qn('r:embed'), rid)

    @staticmethod
    def number_name(ident):
        return 'HN_' + sha(ident.encode())[:24]

    @staticmethod
    def identity(block):
        # Word needs a local bookmark even for anonymous GFM tables. This value
        # never becomes an HRS ID and is meaningful only inside this snapshot.
        return block.get('id') or 'snapshot-table-'+str(block['index'])

    def mark_inline(self, p, name, callback):
        self.bookmark_id += 1
        ident = self.bookmark_id
        p._p.append(element('w:bookmarkStart', id=ident, name=name))
        callback()
        p._p.append(element('w:bookmarkEnd', id=ident))

    def caption(self, b, figure):
        p = self.paragraph('图题' if figure else '表题')
        p.add_run('图 ' if figure else '表 ')
        key=self.identity(b)
        self.mark_inline(p, self.number_name(key), lambda: field(p,
            ' SEQ ' + ('图' if figure else '表') + ' \\* ARABIC ', self.output_numbers[key]))
        text = b['data'].get('caption', '')
        if text:
            p.add_run(' ')
            self.rich(p,text)
        return p

    def data_table(self, b):
        from wps_pipeline.docx_formatting import style_table
        d = b['data']
        self.caption(b, False)
        rows = d['rows']
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                if cell is None: continue
                target = table.cell(ri, ci)
                rs = cell.get('rowspan', 1) if isinstance(cell, dict) else 1
                cs = cell.get('colspan', 1) if isinstance(cell, dict) else 1
                if rs > 1 or cs > 1:
                    target = target.merge(table.cell(ri + rs - 1, ci + cs - 1))
                text = cell['text'] if isinstance(cell, dict) else cell
                for p in list(target._tc.findall(qn('w:p'))): target._tc.remove(p)
                p = target.add_paragraph()
                self.rich(p, text, pre_wrap=True)
        widths = d.get('columnWidths')
        if widths:
            for col, width in zip(table._tbl.tblGrid, widths): col.set(qn('w:w'), str(round(float(width)*100)))
        plain = E.fromstring(E.tostring(table._tbl))
        style_table(plain, self.styles, self.width)
        formatted = parse_xml(E.tostring(plain))
        table._tbl.getparent().replace(table._tbl, formatted)
        for p in formatted.xpath('.//w:p'): zero_indent(p)
        # Explicit per-column source alignment survives template styling; absent
        # values continue to inherit the mother's table paragraph style.
        if d.get('align'):
            for row in formatted.findall(qn('w:tr')):
                column=0
                for cell in row.findall(qn('w:tc')):
                    span=cell.find(qn('w:tcPr')+'/'+qn('w:gridSpan'))
                    choice=d['align'][column] if column<len(d['align']) else None
                    if choice:
                        for p in cell.findall(qn('w:p')):
                            props=p.get_or_add_pPr()
                            for old in list(props):
                                if old.tag==qn('w:jc'):props.remove(old)
                            props.append(element('w:jc',val=choice))
                    column+=int(span.get(qn('w:val'))) if span is not None else 1
        if not b.get('numbered'):
            self.diagnostic('OUTPUT_TABLE_NUMBER_POLICY', b, 'WPS profile numbers this native table; HRS unnumbered source policy is retained in provenance.')

    def group(self, b):
        d = b['data']
        cols = d['columns']
        rows = d.get('rows') or (len(d['images']) + cols - 1) // cols
        # Each logical image row has a separate caption row. A shared row
        # boundary aligns child captions even when the pictures differ in height.
        table = self.doc.add_table(rows=rows * 2, cols=cols)
        table.autofit = False
        pr = table._tbl.tblPr
        for child in list(pr): pr.remove(child)
        pr.extend([element('w:tblW', w=self.width, type='dxa'), element('w:jc', val='center'), element('w:tblLayout', type='fixed'), element('w:tblInd', w=0, type='dxa')])
        borders = element('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'): borders.append(element('w:'+edge, val='nil'))
        pr.append(borders)
        margins = element('w:tblCellMar')
        for edge, value in [('top',65),('bottom',65),('left',85),('right',85)]: margins.append(element('w:'+edge, w=value, type='dxa'))
        pr.append(margins)
        widths = [self.width // cols] * cols
        widths[-1] += self.width - sum(widths)
        for col, width in zip(table._tbl.tblGrid, widths): col.set(qn('w:w'), str(width))
        for row_index in range(rows):
            image_row, caption_row = table.rows[row_index * 2:row_index * 2 + 2]
            # Keep each pair together when it fits, without chaining the whole
            # group or imposing exact heights that would clip a long caption.
            for row in (image_row, caption_row):
                row._tr.get_or_add_trPr().append(element('w:cantSplit'))
            for column, (cell, caption_cell) in enumerate(zip(image_row.cells, caption_row.cells)):
                for target, alignment in ((cell, 'bottom'), (caption_cell, 'top')):
                    target.width = Twips(widths[column])
                    target._tc.get_or_add_tcPr().append(element('w:vAlign', val=alignment))
                    for old in list(target._tc.findall(qn('w:p'))): target._tc.remove(old)
                p = self.paragraph('图形段', cell)
                p.paragraph_format.keep_with_next = True
                caption = self.paragraph('图题', caption_cell)
                caption.paragraph_format.keep_with_next = False
                i = row_index * cols + column
                if i < len(d['images']):
                    img = d['images'][i]
                    self.picture(p, img, widths[column] - 170)
                    if d.get('showSubfigureLabels', True): caption.add_run('(' + chr(97+i) + ')')
                    if img.get('caption'):
                        if d.get('showSubfigureLabels', True): caption.add_run(' ')
                        self.rich(caption,img['caption'])
        self.caption(b, True)

    def space_after_data_table(self, nodes):
        """Use the template's existing caption gap; do not create source/spacer paragraphs."""
        gap=self.doc.styles['表题'].paragraph_format.space_before
        if gap is None or gap<=0:return
        first=next((n for n in nodes if n.tag in (qn('w:p'),qn('w:tbl'))),None)
        if first is None:return
        paragraphs=[first] if first.tag==qn('w:p') else first.xpath('./w:tr[1]/w:tc/w:p[1]')
        for node in paragraphs:
            p=Paragraph(node,self.doc._body)
            before=p.paragraph_format.space_before
            style=p.style;seen=set()
            while before is None and style is not None and style.style_id not in seen:
                seen.add(style.style_id);before=style.paragraph_format.space_before;style=style.base_style
            if (before or 0)<gap:p.paragraph_format.space_before=gap

    def write(self, output):
        self.output_numbers = {}
        # Output sequence is a WPS profile concern, never a replacement HRS numbering rule.
        table_number = 0
        for b in self.ir['blocks']:
            if b['type'] == 'table':
                table_number += 1
                self.output_numbers[self.identity(b)] = table_number
            elif b['type'] in ('figure','image-group'):
                self.output_numbers[b['id']] = b['number']
        self.bookmark_id+=1;scope_id=self.bookmark_id
        scope_name='HBD_'+self.document_key
        self.body.insert(len(self.body)-1,element('w:bookmarkStart',id=scope_id,name=scope_name))
        previous_data_table=False
        for b in self.ir['blocks']:
            self.bookmark_id += 1
            ident = self.bookmark_id
            bookmark = 'HB_' + self.document_key + '_' + str(b['index'])
            begin = element('w:bookmarkStart', id=ident, name=bookmark)
            self.body.insert(len(self.body)-1, begin)
            kind = b['type']
            if kind == 'heading':
                style, heading_text = self.heading_output[b['index']]
                p = self.paragraph(style)
                if b.get('id'):
                    self.mark_inline(p, self.number_name(b['id']), lambda: self.rich(p, heading_text))
                else: self.rich(p, heading_text)
            elif kind == 'markdown':
                for role,value in self.markdown_parts[b['index']]:
                    if role=='markdown':self.markdown(self.model['tokens'][value])
                    else:self.rich(self.paragraph(role),value)
            elif kind == 'code':
                p=self.paragraph('HRS Code')
                # python-docx maps both CR and LF to separate w:br nodes. A
                # source CRLF is one logical line break, never two Word breaks.
                p.add_run(b['body'].replace('\r\n','\n').replace('\r','\n'))
                zero_indent(p._p)
            elif kind == 'equation':
                p = self.paragraph('公式')
                if self.math is None: self.math = MathAdapter(self.rt)
                p._p.append(self.math.omml(b['latex'], True))
                if b.get('numbered'):
                    self.diagnostic('OUTPUT_EQUATION_NUMBER_POLICY', b, 'WPS profile has unnumbered equations; source identity and numbered flag are retained.')
            elif kind == 'table': self.data_table(b)
            elif kind == 'image-group': self.group(b)
            elif kind == 'figure':
                p = self.paragraph('图形段')
                d = b['data']
                if d['kind'] == 'mermaid':
                    self.picture(p, {}, self.width, diagram=b['id'])
                    self.diagnostic('DIAGRAM_RENDITION', b, 'Existing HRS renderer provides a PNG view; original Mermaid source is retained.')
                elif d['kind'] == 'image': self.picture(p, d, self.width)
                else: raise BridgeError('Unsupported figure kind: ' + d['kind'])
                p.paragraph_format.keep_with_next = True
                self.caption(b, True)
            else: raise BridgeError('Unsupported HRS block: ' + kind)
            end=element('w:bookmarkEnd',id=ident);self.body.insert(len(self.body)-1,end)
            nodes = list(self.body)[list(self.body).index(begin)+1:list(self.body).index(end)]
            if previous_data_table:self.space_after_data_table(nodes)
            if nodes:previous_data_table=kind=='table'
            self.metadata_blocks.append({'index': b['index'], 'type':kind, 'id':b.get('id'), 'bookmark':bookmark,
                                         'native_sha256':semantic_hash(nodes, self.doc.part.related_parts,self.doc.part.rels),
                                         'native_drawing_extents':drawing_extents(nodes)})
            if kind == 'heading':
                # Only the output projection is additional metadata; the source
                # level remains in the existing source snapshot/IR.
                self.metadata_blocks[-1]['word_heading_level'] = None if style == 'Title' else int(style.split()[-1])
        self.write_references()
        self.body.insert(len(self.body)-1,element('w:bookmarkEnd',id=scope_id))
        if [E.tostring(node) for node in self.front_nodes]!=self.front_xml:raise BridgeError('Template front matter changed during assembly')
        update=self.doc.settings.element.find(qn('w:updateFields'))
        if update is None:update=element('w:updateFields');self.doc.settings.element.append(update)
        update.set(qn('w:val'),'true')
        assets, blobs = {}, {}
        for path in sorted((self.source/'assets').rglob('*')):
            if path.is_file():
                raw = path.read_bytes()
                digest = sha(raw)
                assets[path.relative_to(self.source).as_posix()] = digest
                blobs[digest] = base64.b64encode(raw).decode('ascii')
        meta = {'schema':'hrs-wps-bridge-v1', 'document_key':self.document_key,'source_sha256':sha(self.raw), 'source_b64':base64.b64encode(self.raw).decode('ascii'),
                'assets':assets,'blobs':blobs,'blocks':self.metadata_blocks,'derived_sections':self.derived_sections,'word_output_options':self.options,
                'body_scope':{'bookmark':scope_name,'kind':'wps-report-body-v1'},'template_sha256':self.template_sha256, 'output_profile':'wps-template-v1','source_render_profile':self.source_profile,'source_render_options':self.source_options,'word_origin':self.word_origin}
        self.doc.save(output)
        self.preserve_template_parts(output)
        provenance_path(output).write_text(json.dumps(meta,ensure_ascii=False,separators=(',',':')),encoding='utf8')
        return {'ok':True, 'output':str(output), 'sha256':sha(output.read_bytes()), 'source_sha256':sha(self.raw),
                'template':self.template_info, 'object_styles':[self.styles[n] for n in ('图形段','图题','表题','表内正文','公式')],
                'diagnostics':self.diagnostics, 'blocks':len(self.metadata_blocks),'word_output_options':self.options}

    def preserve_template_parts(self,output):
        """The original cover/header parts remain byte-exact; only body-owned
        package parts and newly introduced relationships may be regenerated."""
        import zipfile,os
        mutable={'word/document.xml','word/styles.xml','word/numbering.xml','word/settings.xml',
            'word/_rels/document.xml.rels','[Content_Types].xml'}
        if self.note_part is not None:
            name=str(self.note_part.partname).lstrip('/');mutable.add(name)
            part=Path(name);mutable.add(str(part.parent/'_rels'/(part.name+'.rels')).replace('\\','/'))
        with zipfile.ZipFile(self.template) as original,zipfile.ZipFile(output) as generated:
            original_bytes={name:original.read(name) for name in original.namelist() if name not in mutable}
            members={item.filename:(item,generated.read(item.filename)) for item in generated.infolist()}
        temp=output.with_suffix('.preserve.docx')
        try:
            with zipfile.ZipFile(temp,'x',compression=zipfile.ZIP_DEFLATED) as archive:
                for name,(info,data) in members.items():archive.writestr(info,original_bytes.pop(name,data))
                for name,data in original_bytes.items():archive.writestr(name,data)
            os.replace(temp,output)
        finally:
            if temp.exists():temp.unlink()


def prepare(rt, source, *, template=None, accept_numbered_headings=False, options=None):
    return Writer(rt,source,template,accept_numbered_headings,prepare_only=True,options=options).model_request


def convert(rt, source, output, *, template=None, accept_numbered_headings=False, model=None, options=None):
    output = new_path(output)
    sidecar = new_path(provenance_path(output))
    source=Path(source).resolve()
    if source.is_file():source=source.parent
    before={p.relative_to(source).as_posix():sha(p.read_bytes()) for p in source.rglob('*') if p.is_file()}
    writer = Writer(rt, source,template,accept_numbered_headings,model=model,options=options)
    with tempfile.TemporaryDirectory(prefix='bridge-docx-') as temp:
        staged = Path(temp)/'report.docx'
        receipt = writer.write(staged)
        # Publish only a complete, readable OPC package, never half of a conversion.
        Document(staged)
        after={p.relative_to(source).as_posix():sha(p.read_bytes()) for p in source.rglob('*') if p.is_file()}
        if before!=after:raise BridgeError('Source changed during DOCX export')
        if sha(writer.template.read_bytes())!=writer.template_sha256:raise BridgeError('Template changed during DOCX export')
        created=[]
        try:
            for target,raw in ((output,staged.read_bytes()),(sidecar,provenance_path(staged).read_bytes())):
                with target.open('xb') as stream:
                    created.append(target)
                    stream.write(raw)
        except Exception:
            for target in created:target.unlink()
            raise
    receipt.update(output=str(output), sha256=sha(output.read_bytes()),provenance=str(sidecar),provenance_sha256=sha(sidecar.read_bytes()),source_unchanged=True)
    return receipt
